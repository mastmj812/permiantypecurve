"""Dossier zone summary: the per-zone math (risked TC vs Novi, gap flag,
Di nominal + effective, cohort QC) and the summary slide it feeds.
DB-free: a stub TypeCurve and stub Forecast rows."""

from __future__ import annotations

import io
import math
import uuid
from types import SimpleNamespace
from typing import Any

import pytest
from pptx import Presentation
from pptx.shapes.graphfrm import GraphicFrame

from app.db.models import TypeCurve
from app.exports.dossier import SUMMARY_ROWS_PER_SLIDE, build_deal_dossier_pptx
from app.exports.dossier_summary import (
    GAP_FLAG_RATIO,
    SUMMARY_HEADERS,
    CohortQC,
    ZoneStick,
    build_zone_summary,
    cohort_qc,
    novi_eur_per_1000ft,
    one_yr_effective,
    stream_summary,
    summary_cells,
)

_DI_HI = {"oil": 4.0, "gas": 4.0}


def _tc(
    oil_eur: float = 60_000.0, gas_eur: float = 300_000.0, risk: dict[str, Any] | None = None
) -> TypeCurve:
    tc = TypeCurve()
    tc.id = uuid.uuid4()
    tc.name = "alchemist_wcb2"
    tc.included_api10s = ["4230100001", "4230100002", "4230100003"]
    tc.risk_multipliers = risk or {}
    tc.forecast_overrides = {}
    tc.series = {
        "streams": {
            "oil": {"fitted": {"qi": 30.0, "Di": 2.5, "b": 1.0, "eur_per_unit": oil_eur}},
            "gas": {"fitted": {"qi": 90.0, "Di": 1.8, "b": 1.1, "eur_per_unit": gas_eur}},
        }
    }
    return tc


def _fc(qi: float = 500.0, di: float = 2.0, b: float = 1.0, peak: float = 500.0) -> Any:
    return SimpleNamespace(qi=qi, di_initial=di, b=b, peak_rate=peak)


def _stick(cat: str = "PUD", ll: float = 10_000.0, scen: str = "d/s1") -> ZoneStick:
    return ZoneStick(
        well_name="gen-0",
        formation="WCB_2",
        category=cat,
        scenario_ref=scen,
        completed_lateral_ft=ll,
        target_tvd_ft=9_800.0,
    )


def test_effective_matches_house_formula() -> None:
    # nominal 2.5/yr, b=1 -> 1 - 1/(1+2.5) = 71.4% effective
    assert one_yr_effective(2.5, 1.0) == pytest.approx(2.5 / 3.5)
    assert one_yr_effective(1.0, 0.0) == pytest.approx(1 - math.exp(-1.0))
    assert one_yr_effective(None, 1.0) is None


def test_stream_summary_risked_and_gap_flag() -> None:
    s = stream_summary(_tc(risk={"oil": 0.8}), "oil", novi_eur=30_000.0)
    assert s.eur_per_1000ft == pytest.approx(48_000.0)  # 60k x 0.8 — risked as delivered
    assert s.qi_per_1000ft == pytest.approx(24.0)
    assert s.risk_mult == 0.8
    assert s.tc_vs_novi == pytest.approx(0.6)  # 48k / 30k - 1
    assert s.gap_flag  # 1.6x > 1.5x
    ok = stream_summary(_tc(), "oil", novi_eur=50_000.0)
    assert ok.tc_vs_novi == pytest.approx(0.2) and not ok.gap_flag
    low = stream_summary(_tc(), "oil", novi_eur=60_000.0 * GAP_FLAG_RATIO + 1)
    assert low.gap_flag  # flag fires both ways
    none = stream_summary(_tc(), "oil", novi_eur=None)
    assert none.tc_vs_novi is None and not none.gap_flag


def test_novi_eur_matches_the_figure_trapezoid() -> None:
    assert novi_eur_per_1000ft([10.0] * 600) == 6_000.0  # flat: trapezoid == sum
    # declining: trapezoid (the figure's cumFrom), not the rectangle sum
    assert novi_eur_per_1000ft([30.0, 20.0, 10.0]) == pytest.approx(25.0 + 15.0 + 10.0)
    assert novi_eur_per_1000ft(()) is None


def test_cohort_qc_override_wins_and_bounds_per_stream() -> None:
    tc = _tc()
    # well 1: global oil fit Di pinned at the 4.0 cap; well 2: TC override
    # (clean) replaces a pinned global; well 3: no oil forecast at all.
    tc.forecast_overrides = {
        "4230100002": {"oil": {"qi": 400.0, "di_initial": 2.0, "b": 1.0, "peak_rate": 400.0}}
    }
    forecasts = {
        "oil": {"4230100001": _fc(di=3.99), "4230100002": _fc(di=3.99)},
        "gas": {a: _fc() for a in tc.included_api10s},
    }
    qc = cohort_qc(tc, forecasts, _DI_HI)  # type: ignore[arg-type]
    assert qc.n_wells == 3 and qc.n_overridden == 1
    assert qc.at_bound == {"oil": 1, "gas": 0}
    assert qc.di_at_bound == {"oil": 1, "gas": 0} and qc.b_at_bound == {"oil": 0, "gas": 0}
    assert qc.missing == {"oil": 1, "gas": 0}


def test_zone_summary_counts_and_cells() -> None:
    tc = _tc(risk={"oil": 0.8, "gas": 0.8})
    z = build_zone_summary(
        zone_name="WCB_2",
        reserve_category="PUD",
        benches=["WCB_2"],
        tc=tc,
        sticks=[
            _stick("PUD", 9_800.0),
            _stick("UPSIDE", 10_200.0, "d/s2"),
            _stick("PUD", 10_000.0),
        ],
        novi_volumes={"oil": [50.0] * 600, "gas": [500.0] * 600},
        novi_n_sticks=7,
        novi_low_n=False,
        novi_stale=False,
        qc=CohortQC(
            n_wells=3,
            n_overridden=0,
            at_bound={"oil": 2, "gas": 0},
            di_at_bound={"oil": 1, "gas": 0},
            b_at_bound={"oil": 2, "gas": 0},
            missing={"oil": 0, "gas": 0},
        ),
    )
    assert (z.n_sticks, z.n_pud, z.n_upside, z.n_scenarios) == (3, 2, 1, 2)
    assert z.planned_lateral_ft_median == 10_000.0
    assert (
        z.curve_name.startswith("alchemist_wcb2") and z.curve_name != "alchemist_wcb2"
    )  # risked suffix
    cells = summary_cells(z)
    assert len(cells) == len(SUMMARY_HEADERS)
    row = dict(zip(SUMMARY_HEADERS, cells, strict=True))
    assert row["Sticks (PUD / UPSIDE)"] == "3 (2 / 1)"
    assert row["Oil EUR /1,000 ft"] == "48,000" and row["Novi oil"] == "30,000"
    assert row["Oil TC vs Novi"] == "+60% ⚑"
    assert row["Oil Di nom /yr (1-yr eff)"] == "2.50 (71%)"  # nominal with effective beside it
    assert row["Risk"] == "×0.8"
    assert "oil pinned: Di 1 · b 2" in row["QC"]
    assert any("TC 1.60x Novi" in f for f in z.flags)
    assert any("1 cohort oil fits with Di at a bound" in f for f in z.flags)


def _summary(n: int) -> list[Any]:
    tc = _tc()
    return [
        build_zone_summary(
            zone_name=f"Z{i}",
            reserve_category="PUD",
            benches=["WCB_2"],
            tc=tc,
            sticks=[_stick()],
            novi_volumes=None,
            novi_n_sticks=0,
            novi_low_n=False,
            novi_stale=False,
            qc=CohortQC(n_wells=3, n_overridden=0),
            novi_error="Novi comparison unavailable: down",
        )
        for i in range(n)
    ]


def test_summary_slide_leads_the_deck() -> None:
    content = build_deal_dossier_pptx(SimpleNamespace(), [], [], summary=_summary(2))  # type: ignore[arg-type]
    pres = Presentation(io.BytesIO(content))
    assert len(pres.slides) == 1
    tables = [s for s in pres.slides[0].shapes if isinstance(s, GraphicFrame) and s.has_table]
    assert len(tables) == 1
    t = tables[0].table
    assert len(t.rows) == 3 and len(t.columns) == len(SUMMARY_HEADERS)
    assert t.cell(0, 0).text == "Zone" and t.cell(1, 0).text == "Z0"
    assert "Novi unavailable" in t.cell(1, len(SUMMARY_HEADERS) - 1).text


def test_summary_paginates() -> None:
    content = build_deal_dossier_pptx(
        SimpleNamespace(),
        [],
        [],
        summary=_summary(SUMMARY_ROWS_PER_SLIDE + 1),  # type: ignore[arg-type]
    )
    assert len(Presentation(io.BytesIO(content)).slides) == 2


def test_parse_linestring_wkt() -> None:
    from app.exports.dossier_summary import parse_linestring_wkt

    assert parse_linestring_wkt("LINESTRING(-103.5 31.6, -103.49 31.62)") == (
        (-103.5, 31.6),
        (-103.49, 31.62),
    )
    assert parse_linestring_wkt(None) == ()
    assert parse_linestring_wkt("POINT(1 2)") == ()


def test_overview_and_support_slides_follow_the_summary() -> None:
    from pptx.shapes.picture import Picture

    from app.exports.dossier import ComparisonSlideInput, ScenarioSlideInput
    from tests.test_dossier_export import _PNG

    content = build_deal_dossier_pptx(
        SimpleNamespace(),  # type: ignore[arg-type]
        [ScenarioSlideInput(title="plan_a", subtitle="", map_png=_PNG, gunbarrel_png=_PNG)],
        [],
        summary=_summary(1),
        overview=ComparisonSlideInput(
            title="Curve assignment overview", subtitle="", figure_png=_PNG
        ),
        supports=[
            ScenarioSlideInput(
                title="WCB_2 — 2 sticks take wcb2, built from 18 wells",
                subtitle="TC oil 60.6 bbl/ft",
                map_png=_PNG,
                gunbarrel_png=_PNG,
            )
        ],
    )
    pres = Presentation(io.BytesIO(content))
    titles = [
        " ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in pres.slides
    ]
    assert len(pres.slides) == 4
    assert "Zone summary" in titles[0]
    assert "Curve assignment overview" in titles[1]
    assert "2 sticks take wcb2" in titles[2]
    assert "plan_a" in titles[3]
    assert len([s for s in pres.slides[2].shapes if isinstance(s, Picture)]) == 2  # map | zoom


def test_long_title_shrinks_and_long_subtitle_wraps() -> None:
    from pptx.util import Pt

    from app.exports.dossier import ComparisonSlideInput
    from tests.test_dossier_export import _PNG

    title = "BS1_S — Type Curve vs Novi ML (n=20: 20 PUD / 0 RES)"  # 52 chars: wrapped before
    content = build_deal_dossier_pptx(
        SimpleNamespace(),  # type: ignore[arg-type]
        [],
        [],
        summary=_summary(1),
        overview=ComparisonSlideInput(title=title, subtitle="x" * 250, figure_png=_PNG),
    )
    slide = Presentation(io.BytesIO(content)).slides[1]
    boxes = [sh for sh in slide.shapes if sh.has_text_frame]
    title_runs = [
        r
        for sh in boxes
        if sh.text_frame.text == title
        for p in sh.text_frame.paragraphs
        for r in p.runs
    ]
    assert title_runs and all(r.font.size == Pt(24) for r in title_runs)
    sub = next(sh for sh in boxes if sh.text_frame.text.startswith("xxx"))
    assert sub.text_frame.word_wrap is True
    pic = next(sh for sh in slide.shapes if sh.shape_type == 13)  # picture
    assert (pic.top + pic.height) / 914400 <= 6.85  # clear of the footer band


def _cw(api: str, lon: float, **k: Any) -> Any:
    from app.exports.dossier_summary import CohortWell

    return CohortWell(
        api10=api,
        name=f"W{api}",
        lateral_ft=10_000.0,
        oil_eur_per_ft=60.0,
        coords=((lon, 31.6), (lon, 31.62)),
        **k,
    )


def test_came_on_and_pinned_tags() -> None:
    from app.exports.dossier_summary import came_on, pinned_params

    assert came_on(_cw("1", 0, scenario_class="topfill", parents_below=("WCA_1",))) == "over WCA_1"
    assert (
        came_on(_cw("1", 0, scenario_class="underfill", parents_above=("BS3_C",))) == "under BS3_C"
    )
    assert (
        came_on(_cw("1", 0, scenario_class="sandwich", parents_above=("A",), parents_below=("B",)))
        == "between A / B"
    )
    assert (
        came_on(_cw("1", 0, scenario_class="codev_stack", codev_benches=("BS2_C",)))
        == "co-developed with BS2_C"
    )
    assert came_on(_cw("1", 0, scenario_class="standalone")) == "alone"
    assert came_on(_cw("1", 0)) == "—"  # not in dev_scenario
    assert pinned_params("Di at upper bound (4.0); b at lower bound (0.9)") == ("Di hi", "b lo")
    assert pinned_params(None) == ()


def test_cohort_rows_nearest_first_with_fit_source() -> None:
    from app.exports.dossier_summary import COHORT_HEADERS, cohort_rows

    stick = ZoneStick(
        well_name="s",
        formation="WCB_2",
        category="PUD",
        scenario_ref="d/s",
        completed_lateral_ft=10_000.0,
        target_tvd_ft=None,
        legs_lonlat=((-103.5, 31.6, -103.5, 31.62),),
    )
    far = _cw("2", -103.3, fit_source="override", pinned=("Di hi",))
    near = _cw("1", -103.51, fit_source="global", fit_edited=True)
    nogeo = _cw("3", 0.0, fit_source="none")
    nogeo = type(nogeo)(**{**nogeo.__dict__, "coords": ()})
    rows = cohort_rows([far, nogeo, near], [stick])
    col = {h: i for i, h in enumerate(COHORT_HEADERS)}
    assert [r[col["api10"]] for r in rows] == ["1", "2", "3"]  # nearest first, no geometry last
    assert rows[0][col["Oil fit"]] == "global (edited)" and rows[1][col["Oil fit"]] == "TC override"
    assert rows[1][col["Oil fit pinned"]] == "Di hi" and rows[2][col["To nearest stick mi"]] == "—"
    assert float(rows[0][col["To nearest stick mi"]]) == pytest.approx(
        0.6, abs=0.05
    )  # 0.01 deg lon at 31.6N


def test_cohort_table_precedes_stream_slides_and_paginates() -> None:
    from app.exports.dossier import COHORT_ROWS_PER_SLIDE, CohortTableInput
    from tests.test_deal_export import _curve
    from tests.test_dossier_export import _curve_input, _StubSession

    tc = _curve("holdTheLine_wca_v1")
    rows = [
        [
            f"W{i}",
            "op",
            str(i),
            "2024-01",
            "10,000",
            "1.0",
            "60.0",
            "global",
            "Di hi" if i == 0 else "—",
            "alone",
        ]
        for i in range(COHORT_ROWS_PER_SLIDE + 3)
    ]
    content = build_deal_dossier_pptx(
        _StubSession([tc]),  # type: ignore[arg-type]
        [],
        [_curve_input(tc)],
        cohort_tables={tc.id: CohortTableInput(curve_name=tc.name, rows=rows)},
    )
    pres = Presentation(io.BytesIO(content))
    titles = [
        " ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame) for sl in pres.slides
    ]
    assert len(pres.slides) == 5  # 2 table pages + oil/gas/water
    assert "curve wells (19) 1/2" in titles[0] and "curve wells (19) 2/2" in titles[1]
    assert "holdTheLine_wca_v1 Oil" in titles[2]
    t = next(s for s in pres.slides[0].shapes if isinstance(s, GraphicFrame) and s.has_table).table
    assert len(t.rows) == 10 + 1  # 19 wells -> balanced 10 + 9
