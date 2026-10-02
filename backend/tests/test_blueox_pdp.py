"""Blue Ox PDP workbook (contract §2) — pure builder + volume rule, DB-free."""

from __future__ import annotations

import io
from datetime import date

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from app.exports.blueox import BlueOxContractError
from app.exports.blueox_pdp import (
    DOLLAR_BASIS_NA,
    MODE_CURVES,
    PdpExportData,
    PdpGroup,
    PdpWell,
    build_pdp_workbook,
    first_row_month,
    pdp_filename,
    readiness,
)
from app.forecasting import daily
from app.pdp.export import delivered_daily, lease_of, month_starts, monthly_vector

N = 24


def _well(api10: str, name: str, scale: float = 1.0, **kw: object) -> PdpWell:
    vols = {
        s: [scale * k * (N - i) for i in range(N)]
        for s, k in (("oil", 100.0), ("gas", 600.0), ("water", 800.0))
    }
    return PdpWell(
        api10=api10,
        well_name=name,
        volumes=vols,
        methods=kw.get("methods", {"oil": "daily_fit", "gas": "daily_fit", "water": "daily_fit"}),  # type: ignore[arg-type]
        locked=kw.get("locked", {"oil": True, "gas": True, "water": True}),  # type: ignore[arg-type]
        review_flags=kw.get("flags", ()),  # type: ignore[arg-type]
    )


def _data(groups: list[PdpGroup], **kw: object) -> PdpExportData:
    eff = kw.get("effective_date", date(2026, 8, 1))
    return PdpExportData(
        codename="alchemist",
        export_date=date(2026, 10, 2),
        effective_date=eff,  # type: ignore[arg-type]
        first_row_month=first_row_month(eff),  # type: ignore[arg-type]
        curve_months=N,
        groups=groups,
        production_history_through=date(2026, 7, 3),
        source_system="test",
        curve_params_source="test",
        prepared_by="mast",
        vdr_id="alchemist_merlin",
        zone_names=kw.get("zones", ()),  # type: ignore[arg-type]
    )


def test_filename_and_first_row_month() -> None:
    assert pdp_filename("alchemist", date(2026, 10, 2)) == "alchemist_pdp_2026-10-02.xlsx"
    assert first_row_month(date(2026, 8, 1)) == date(2026, 8, 1)
    assert first_row_month(date(2026, 8, 15)) == date(2026, 9, 1)
    assert first_row_month(date(2026, 12, 31)) == date(2027, 1, 1)


def test_workbook_shape_and_manifest_ties_to_sheets() -> None:
    a, b = _well("4230135457", "Lowe 74 Unit 61H"), _well("4230135458", "Lowe 74 Unit 63H", 2.0)
    wb = load_workbook(io.BytesIO(build_pdp_workbook(_data([PdpGroup("Lowe 74 Unit", [a, b])]))))
    assert wb.sheetnames == ["Lowe 74 Unit", "manifest"]
    ws = wb["Lowe 74 Unit"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("gross_oil_bbl", "gross_gas_mcf", "gross_water_bbl")  # no date column
    assert len(rows) == N + 1
    assert rows[1][0] == pytest.approx(3 * 100.0 * N)  # aggregate of both wells, month 1
    sheet_sums = [sum(r[c] for r in rows[1:]) for c in range(3)]

    man = list(wb["manifest"].iter_rows(values_only=True))
    kv = {r[0]: r[1] for r in man if r and r[0] and r[1] is not None and len(r) >= 2}
    assert kv["deliverable"] == "pdp"
    assert kv["first_row_month"] == "2026-08"
    assert kv["water_basis"] == "seller_measured_daily_vdr"
    hdr = next(i for i, r in enumerate(man) if r and r[0] == "group")
    grp = man[hdr + 1]
    assert (
        grp[0] == "Lowe 74 Unit"
        and grp[1] == MODE_CURVES
        and grp[2] == 2
        and grp[3] == DOLLAR_BASIS_NA
    )
    assert list(grp[4:7]) == pytest.approx(sheet_sums, rel=0, abs=1e-6)  # exact tie


@pytest.mark.parametrize(
    "groups, zones, needle",
    [
        ([PdpGroup("WCB_2 West", [_well("1", "a")])], ("WCB_2 West",), "zone name"),
        (
            [PdpGroup("g1", [_well("1", "a")]), PdpGroup("G1", [_well("2", "b")])],
            (),
            "duplicate group",
        ),
        ([PdpGroup("g1", [_well("1", "a")]), PdpGroup("g2", [_well("1", "a")])], (), "is in both"),
        ([PdpGroup("x" * 32, [_well("1", "a")])], (), "longer than"),
        ([PdpGroup("a/b", [_well("1", "a")])], (), "contains"),
        ([PdpGroup("manifest", [_well("1", "a")])], (), "reserved"),
        (
            [
                PdpGroup(
                    "g",
                    [_well("1", "a", methods={"oil": "none", "gas": "daily_fit", "water": "none"})],
                )
            ],
            (),
            "no oil forecast",
        ),
    ],
)
def test_contract_violations_block(
    groups: list[PdpGroup], zones: tuple[str, ...], needle: str
) -> None:
    with pytest.raises(BlueOxContractError, match=needle):
        build_pdp_workbook(_data(groups, zones=zones))


def test_vector_length_mismatch_blocks() -> None:
    w = _well("1", "a")
    w.volumes["gas"].pop()
    with pytest.raises(BlueOxContractError, match="length"):
        build_pdp_workbook(_data([PdpGroup("g", [w])]))


def test_readiness_lists_unlocked_missing_and_flags() -> None:
    w = _well(
        "1",
        "Sheridan 27A Unit 73H",
        locked={"oil": True, "gas": False, "water": False},
        methods={"oil": "daily_fit", "gas": "daily_fit", "water": "none"},
        flags=("at_bound", "recent_break"),
    )
    found = {(f.check, f.detail) for f in readiness(_data([PdpGroup("g", [w])]))}
    assert ("unlocked", "Sheridan 27A Unit 73H: gas not locked") in found
    assert ("no_forecast", "Sheridan 27A Unit 73H: no water forecast (zeros)") in found
    assert ("review_flags", "Sheridan 27A Unit 73H: recent_break") in found  # at_bound not echoed


def test_delivered_daily_actuals_then_model_then_zero() -> None:
    days = pd.date_range("2026-06-01", "2026-07-03", freq="D")
    actual = pd.Series(100.0, index=days)
    actual.iloc[5] = np.nan  # no report -> 0
    params = {"qi": 90.0, "Di": 0.5, "b": 1.0, "Df": 0.08}
    out = delivered_daily(
        actual,
        params=params,
        anchor=date(2026, 7, 3),
        data_through=date(2026, 7, 3),
        uptime=0.9,
        start=date(2026, 7, 1),
        end=date(2026, 8, 31),
        horizon_end=date(2026, 8, 15),
    )
    assert out.loc["2026-07-01":"2026-07-03"].tolist() == [100.0, 100.0, 100.0]
    model_jul4 = float(
        daily.model_rate(params, date(2026, 7, 3), pd.DatetimeIndex(["2026-07-04"]))[0]
    )
    assert out.loc["2026-07-04"] == pytest.approx(model_jul4 * 0.9)
    assert (out.loc["2026-08-16":] == 0).all()
    months = month_starts(date(2026, 7, 1), 2)
    vec = monthly_vector(out, months)
    assert sum(vec) == pytest.approx(out.sum())


def test_delivered_daily_without_forecast_is_history_only() -> None:
    days = pd.date_range("2026-07-01", "2026-07-03", freq="D")
    out = delivered_daily(
        pd.Series(5.0, index=days),
        params=None,
        anchor=None,
        data_through=date(2026, 7, 3),
        uptime=1.0,
        start=date(2026, 7, 1),
        end=date(2026, 7, 31),
        horizon_end=date(2076, 7, 1),
    )
    assert out.sum() == pytest.approx(15.0)


def test_month_starts_and_lease() -> None:
    assert month_starts(date(2026, 11, 1), 3) == [
        date(2026, 11, 1),
        date(2026, 12, 1),
        date(2027, 1, 1),
    ]
    assert lease_of("Sheridan 27A Unit 73H") == "Sheridan 27A Unit"
    assert lease_of("Solo") == "Solo"
