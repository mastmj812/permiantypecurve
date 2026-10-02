"""Blue Ox Deliverable B — the PDP workbook (contract §2), pure builder.

``<codename>_pdp_<YYYY-MM-DD>.xlsx``: one sheet per PDP group (gross
aggregate monthly volumes, NO date column, row 1 = ``first_row_month``)
plus a ``manifest`` sheet. Curves mode only — this stack never emits
revenue/opex (scope rule); ``dollar_basis`` is declared
``not_applicable``.

No DB / HTTP here (same pattern as ``exports/blueox.py``): the service
(``app.pdp.export``) assembles :class:`PdpExportData` from
``pdp_forecasts`` + the synced daily rows; this module validates against
the contract and writes bytes.

Conventions declared in the manifest (ledger §12, pending Blue Ox ack):

* ``first_row_month`` — the month row 1 covers: the effective-date month
  when the effective date is the 1st, else the next month.
* ``history_basis`` — days on or before ``production_history_through``
  carry the seller's REPORTED daily volumes; later days are the PDP
  forecast (producing-day model rate x per-well uptime).
* ``water_basis`` — seller-measured daily water (VDR), not Novi's TX
  allocation.
* Volumes stop at each well's first production + 50 years (raw technical
  horizon, no economic limit); later rows are zero.
* NGL: no ``gross_ngl_bbl`` column — Blue Ox applies its own yield
  (2026-07-20 NGL amendment, same basis as the curve drop).
"""

from __future__ import annotations

import io
import re
from collections.abc import Sequence
from dataclasses import dataclass, field
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.exports.blueox import GAS_BASIS, NGL_BASIS, RISKING_UNRISKED, BlueOxContractError

PDP_COLUMNS: tuple[tuple[str, str], ...] = (
    ("gross_oil_bbl", "oil"),
    ("gross_gas_mcf", "gas"),
    ("gross_water_bbl", "water"),
)
PDP_STREAMS: tuple[str, ...] = ("oil", "gas", "water")
SHEET_NAME_MAX = 31
_FORBIDDEN_RE = re.compile(r"[:\\/?*\[\]]")
RESERVED_PDP_SHEETS: frozenset[str] = frozenset({"manifest"})
MODE_CURVES = "curves"
DOLLAR_BASIS_NA = "not_applicable"
HISTORY_BASIS = "seller_reported_daily_through_history_then_forecast"
WATER_BASIS = "seller_measured_daily_vdr"
FORECAST_METHOD = (
    "producing-day modified hyperbolic on seller daily volumes (log residuals, t=0 at each "
    "stream's own peak) x per-well uptime (routine downtime, trailing 365 d); unpeaked wells "
    "decline from the trailing-30-d rate on same-bench cohort Di"
)
HORIZON_NOTE = "each well first production + 50 yr (raw technical, no economic limit)"

_BOLD = Font(bold=True)


@dataclass(frozen=True)
class PdpWell:
    """One well's monthly volumes on the workbook grid + review status."""

    api10: str
    well_name: str
    volumes: dict[str, list[float]]  # stream -> curve_months values
    methods: dict[str, str]  # stream -> daily_fit | transfer_now | manual | none
    locked: dict[str, bool]
    review_flags: Sequence[str] = ()
    uptime_factor: float = 1.0


@dataclass(frozen=True)
class PdpGroup:
    name: str
    wells: Sequence[PdpWell]


@dataclass
class PdpExportData:
    codename: str
    export_date: date
    effective_date: date
    first_row_month: date
    curve_months: int
    groups: Sequence[PdpGroup]
    production_history_through: date
    source_system: str
    curve_params_source: str
    prepared_by: str
    vdr_id: str
    zone_names: Sequence[str] = ()
    supersedes: str | None = None
    extra_manifest: list[tuple[str, Any]] = field(default_factory=list)


@dataclass(frozen=True)
class PdpFinding:
    status: str  # warn | fail
    check: str
    detail: str


def pdp_filename(codename: str, export_date: date) -> str:
    return f"{codename}_pdp_{export_date.isoformat()}.xlsx"


def first_row_month(effective_date: date) -> date:
    """Row 1 = the effective-date month when the date is the 1st, else the
    next month (contract §2: 'the first month after the effective date')."""
    if effective_date.day == 1:
        return effective_date
    y, m = (
        (effective_date.year + 1, 1)
        if effective_date.month == 12
        else (effective_date.year, effective_date.month + 1)
    )
    return date(y, m, 1)


def group_totals(g: PdpGroup, curve_months: int) -> dict[str, list[float]]:
    out = {s: [0.0] * curve_months for s in PDP_STREAMS}
    for w in g.wells:
        for s in PDP_STREAMS:
            for i, v in enumerate(w.volumes[s]):
                out[s][i] += float(v)
    return out


def _sheet_name_problem(name: str) -> str | None:
    if not name or name != name.strip():
        return "empty or leading/trailing spaces"
    if len(name) > SHEET_NAME_MAX:
        return f"longer than {SHEET_NAME_MAX} characters"
    if _FORBIDDEN_RE.search(name) or name.startswith("'") or name.endswith("'"):
        return "contains : \\ / ? * [ ] or a leading/trailing apostrophe"
    if name.lower() in RESERVED_PDP_SHEETS:
        return "reserved sheet name"
    return None


def _validate(data: PdpExportData) -> None:
    errs: list[str] = []
    if data.curve_months <= 0:
        errs.append("curve_months must be positive")
    if data.first_row_month != first_row_month(data.effective_date):
        errs.append("first_row_month does not follow from effective_date")
    if not data.groups:
        errs.append("no PDP groups")
    zones = {z.lower() for z in data.zone_names}
    seen_groups: set[str] = set()
    seen_wells: dict[str, str] = {}
    for g in data.groups:
        problem = _sheet_name_problem(g.name)
        if problem:
            errs.append(f"group {g.name!r}: {problem}")
        if g.name.lower() in zones:
            errs.append(f"group {g.name!r} equals a curve-drop zone name (contract §2)")
        if g.name.lower() in seen_groups:
            errs.append(f"duplicate group name {g.name!r} (case-insensitive)")
        seen_groups.add(g.name.lower())
        if not g.wells:
            errs.append(f"group {g.name!r} has no wells")
        for w in g.wells:
            if w.api10 in seen_wells:
                errs.append(f"well {w.api10} is in both {seen_wells[w.api10]!r} and {g.name!r}")
            seen_wells[w.api10] = g.name
            if w.methods.get("oil", "none") == "none":
                errs.append(f"well {w.api10} ({w.well_name}) has no oil forecast")
            for s in PDP_STREAMS:
                vec = w.volumes.get(s)
                if vec is None or len(vec) != data.curve_months:
                    errs.append(f"well {w.api10} {s}: vector length != curve_months")
                elif any(v != v for v in vec):  # NaN
                    errs.append(f"well {w.api10} {s}: NaN in vector")
        totals = group_totals(g, data.curve_months) if g.wells else None
        if totals:
            for s in PDP_STREAMS:
                if any(v < 0 for v in totals[s]):
                    errs.append(f"group {g.name!r} {s}: negative monthly volume")
    if errs:
        raise BlueOxContractError("PDP workbook contract violations:\n- " + "\n- ".join(errs))


def readiness(data: PdpExportData) -> tuple[PdpFinding, ...]:
    """Pre-send review status (warnings — the engineer decides). Every
    unlocked stream and every open review flag is listed."""
    out: list[PdpFinding] = []
    for g in data.groups:
        for w in g.wells:
            unlocked = [
                s for s in PDP_STREAMS if w.methods.get(s, "none") != "none" and not w.locked.get(s)
            ]
            if unlocked:
                out.append(
                    PdpFinding(
                        "warn", "unlocked", f"{w.well_name}: {', '.join(unlocked)} not locked"
                    )
                )
            missing = [s for s in PDP_STREAMS if w.methods.get(s, "none") == "none"]
            if missing:
                out.append(
                    PdpFinding(
                        "warn",
                        "no_forecast",
                        f"{w.well_name}: no {', '.join(missing)} forecast (zeros)",
                    )
                )
            flags = sorted({f for f in w.review_flags if f != "at_bound"})
            if flags:
                out.append(PdpFinding("warn", "review_flags", f"{w.well_name}: {', '.join(flags)}"))
    return tuple(out)


def _write_group(ws: Any, g: PdpGroup, curve_months: int) -> None:
    totals = group_totals(g, curve_months)
    ws.append([h for h, _ in PDP_COLUMNS])
    for c in range(1, len(PDP_COLUMNS) + 1):
        ws.cell(row=1, column=c).font = _BOLD
    for i in range(curve_months):
        ws.append([round(totals[s][i], 4) for _, s in PDP_COLUMNS])
    for c in range(1, len(PDP_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18


def _write_manifest(ws: Any, data: PdpExportData) -> None:
    n_wells = sum(len(g.wells) for g in data.groups)
    block_a: list[tuple[str, Any]] = [
        ("deal_codename", data.codename),
        ("deliverable", "pdp"),
        ("export_date", data.export_date.isoformat()),
        ("effective_date", data.effective_date.isoformat()),
        ("first_row_month", data.first_row_month.strftime("%Y-%m")),
        ("source_system", data.source_system),
        (
            "governing_export",
            f"{pdp_filename(data.codename, data.export_date)} (all {n_wells} PDP wells)",
        ),
        ("gas_basis", GAS_BASIS),
        ("ngl_basis", NGL_BASIS),
        ("water_basis", WATER_BASIS),
        ("risking", RISKING_UNRISKED),
        ("curve_months", data.curve_months),
        ("production_history_through", data.production_history_through.isoformat()),
        ("history_basis", HISTORY_BASIS),
        ("forecast_method", FORECAST_METHOD),
        ("forecast_horizon", HORIZON_NOTE),
        ("curve_params_source", data.curve_params_source),
        ("vdr_id", data.vdr_id),
        ("prepared_by", data.prepared_by),
    ]
    if data.supersedes:
        block_a.append(("supersedes", data.supersedes))
    block_a.extend(data.extra_manifest)
    for k, v in block_a:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).font = _BOLD
    ws.append([])
    ws.append([])

    # Block B — one row per group, computed from the same vectors the
    # group sheets were written from (ties to the sheets exactly).
    header = [
        "group",
        "mode",
        "well_count",
        "dollar_basis",
        "eur_oil_bbl",
        "eur_gas_mcf",
        "eur_water_bbl",
        "api10s",
    ]
    ws.append(header)
    for c in range(1, len(header) + 1):
        ws.cell(row=ws.max_row, column=c).font = _BOLD
    for g in data.groups:
        t = group_totals(g, data.curve_months)
        ws.append(
            [
                g.name,
                MODE_CURVES,
                len(g.wells),
                DOLLAR_BASIS_NA,
                sum(round(v, 4) for v in t["oil"]),
                sum(round(v, 4) for v in t["gas"]),
                sum(round(v, 4) for v in t["water"]),
                ", ".join(w.api10 for w in g.wells),
            ]
        )
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60
    for c in range(3, len(header) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 18


def build_pdp_workbook(data: PdpExportData) -> bytes:
    """Validate (raises :class:`BlueOxContractError`) and emit the bytes.
    Readiness warnings are NOT blocking — callers surface them."""
    _validate(data)
    wb = Workbook()
    wb.remove(wb.active)
    for g in data.groups:
        _write_group(wb.create_sheet(g.name), g, data.curve_months)
    _write_manifest(wb.create_sheet("manifest"), data)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
